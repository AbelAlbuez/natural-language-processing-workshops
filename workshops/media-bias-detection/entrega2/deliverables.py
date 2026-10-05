import base64
import json
import os
from pathlib import Path

import nbformat
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


LIMITATIONS = [
    "Todos los titulos disponibles son proxies slug; no titulares editoriales verificados.",
    "No hay cuerpos HTML: no se pueden estudiar citas, actores en contexto o framing textual completo.",
    "Cinco medios principales y Cambio con 98 articulos: no ocho medios comparables.",
    "Las fechas month y lastmod requieren cautela; un corte por fecha no certifica gobierno.",
    "Keywords seleccionan temas/candidatos, no eventos verificados ni omisiones demostradas.",
    "LDA/coherencia se evaluan en entrenamiento; el mejor K no es un optimo universal.",
    "Programas, emisiones y audio forman parte del archivo; la mezcla de formatos puede dominar las frecuencias.",
]


def latex_escape(text):
    mapping = {"\\": r"\textbackslash{}", "_": r"\_", "%": r"\%", "&": r"\&", "#": r"\#", "$": r"\$", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(mapping.get(character, character) for character in str(text))


def number(value):
    return "no definido" if value is None else f"{value:.3f}"


def generate(output, report_dir):
    report_dir.mkdir(parents=True, exist_ok=True)
    summary = json.loads((output / "resumen.json").read_text())
    if "lda" not in summary:
        raise ValueError("LDA no ha terminado; no generar conclusiones finales")
    lda = summary["lda"]
    lexical = summary["textometry"]
    coverage = pd.read_csv(output / "cobertura_corpus.csv")
    optimization = pd.read_csv(output / "lda_optimization.csv")
    topics = pd.read_csv(output / "lda_topics.csv")
    events = pd.read_csv(output / "eventos_polemicos.csv")
    rates = events.pivot(index="source_id", columns="tema", values="pct_medio")
    ngram_leader = pd.read_csv(output / "bigrams_top100.csv").iloc[0]
    event_findings = (
        f"La Republica tiene {rates.loc['la_republica', 'Reforma tributaria']:.2f}% de coincidencias tributarias "
        f"respecto a sus titulos; RCN y Blu registran {rates.loc['noticias_rcn', 'Conflicto armado']:.2f}% "
        f"y {rates.loc['blu_radio', 'Conflicto armado']:.2f}% en el tema amplio de conflicto. "
        "Son proporciones de keywords, no medidas validadas de framing."
    )
    format_finding = (
        f"El bigrama mas frecuente es '{ngram_leader['ngram']}' ({int(ngram_leader['frecuencia']):,} apariciones). "
        "Programacion de radio y emisiones condicionan los resultados; no confundir formato con posicion editorial."
    )
    metrics = pd.DataFrame([
        ("Articulos del export", summary["articulos"]), ("Medios presentes", summary["medios"]),
        ("Titulos disponibles (slug)", summary["con_titulo"]), ("Titulos ausentes", summary["sin_titulo"]),
        ("Vocabulario depurado", lexical["vocabulario"]), ("Documentos utiles LDA", lda["documentos_lda"]),
        ("K con mayor coherencia c_v", lda["best_k"]), ("Coherencia c_v", lda["coherence_cv"]),
    ], columns=["Metrica", "Valor"])
    metrics.to_csv(output / "resumen_analisis.csv", index=False)
    narrative = (
        f"El export contiene {summary['articulos']:,} registros de {summary['medios']} medios entre "
        f"{summary['desde']} y {summary['hasta']}. Hay {summary['con_titulo']:,} titulos slug y "
        f"{summary['sin_titulo']:,} ausentes. Se entrenaron K={lda['topics_grid']} sobre "
        f"{lda['documentos_lda']:,} documentos utiles, con {lda['passes']} pasadas y semilla 42. "
        f"K={lda['best_k']} obtuvo el mayor c_v ({lda['coherence_cv']:.4f}) de la grilla."
    )
    notebook = nbformat.v4.new_notebook()
    notebook.metadata = {"kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"}, "language_info": {"name": "python"}}
    cells = []
    def markdown(text):
        cell = nbformat.v4.new_markdown_cell(text)
        cell.metadata["language"] = "markdown"
        cell.metadata["id"] = cell.id
        cells.append(cell)
    def code(source):
        cell = nbformat.v4.new_code_cell(source)
        cell.metadata["language"] = "python"
        cell.metadata["id"] = cell.id
        cells.append(cell)
    def image(filename, caption):
        markdown(caption)
        code(f"display(Image(filename=str(RESULTS / 'figuras' / '{filename}')))" )
    markdown("# Entrega 2: corpus Duque, textometria y topicos\n\n" + narrative)
    markdown("## 1. Alcance y limitaciones\n\n" + "\n".join("- " + item for item in LIMITATIONS))
    code("from pathlib import Path\nimport json\nimport pandas as pd\nfrom IPython.display import Image, display\ncandidates = [Path('resultados'), Path('entrega2/resultados'), Path('workshops/media-bias-detection/entrega2/resultados'), Path('natural-language-processing-workshops/workshops/media-bias-detection/entrega2/resultados')]\nRESULTS = next((path for path in candidates if (path / 'resumen.json').exists()), None)\nassert RESULTS is not None, 'No se encuentra la carpeta de resultados de Entrega 2'\nsummary = json.loads((RESULTS / 'resumen.json').read_text())\ndisplay(pd.read_csv(RESULTS / 'resumen_analisis.csv'))")
    code("display(pd.read_csv(RESULTS / 'cobertura_corpus.csv'))\ndisplay(pd.read_csv(RESULTS / 'precision_gobierno.csv'))")
    image("cobertura_corpus.png", "La representacion de medios es desigual; Cambio es una submuestra tardia.")
    markdown("## 2. Textometria\n\nMinusculas y tildes normalizadas; stopwords NLTK y verbos de reporte. Se conservan geografia y actores. Una pendiente/correlacion log-log es descriptiva, no una prueba de ley de potencia. TTR depende de longitud; MSTTR agregado usa ventanas completas de 50 tokens. Yule K = 10000 * sum(f*(f-1))/N^2. Los n-gramas son adyacencias despues de filtrar stopwords.")
    code("display(pd.read_csv(RESULTS / 'frecuencias.csv').head(20))\ndisplay(pd.read_csv(RESULTS / 'diversidad_por_medio.csv'))")
    image("zipf_law.png", f"Ajuste descriptivo: pendiente {number(lexical['zipf_slope'])}, R2={number(lexical['zipf_r2'])}; no certifica Zipf.")
    image("lexical_diversity.png", "Diversidad por medio: no comparar TTR agregado sin controlar longitud/muestra.")
    image("ngrams.png", "N-gramas depurados; frecuencia no equivale a significancia estadistica. " + format_finding)
    code("display(pd.read_csv(RESULTS / 'bigrams_top100.csv').head(30))\ndisplay(pd.read_csv(RESULTS / 'trigrams_top50.csv').head(20))")
    markdown("## 3. Temas y casos temporales\n\nProporciones respecto a los titulos disponibles de cada medio. No hallar keywords no prueba ausencia de cobertura. Los casos temporales son candidatos a revision humana, no clusters certificados del mismo hecho. Los log-odds son descriptivos, sin significancia inferencial.")
    code("display(pd.read_csv(RESULTS / 'eventos_polemicos.csv'))\ndisplay(pd.read_csv(RESULTS / 'casos_temporales.csv'))\ndisplay(pd.read_csv(RESULTS / 'ejemplos_casos.csv').head(18))\ndisplay(pd.read_csv(RESULTS / 'contrastes_lexicos.csv').head(30))")
    image("eventos_framing.png", "Cobertura relativa de temas, no prueba de framing o bias. " + event_findings)
    markdown("## 4. LDA sobre corpus completo\n\nSe excluyen documentos vacios tras el filtrado del diccionario y se conservan con topico -1 en la asignacion. K se selecciona por c_v en entrenamiento. Perplejidad = 2**(-log_bound); no se confunde con el bound negativo.")
    code("display(pd.read_csv(RESULTS / 'lda_optimization.csv'))\ndisplay(pd.read_csv(RESULTS / 'lda_topics.csv'))\nassignments = pd.read_parquet(RESULTS / 'document_topics.parquet')\nassert assignments['article_id'].is_unique\nassert len(assignments) == summary['articulos']\ndisplay(assignments.head())")
    image("lda_optimization.png", f"Mejor K en la grilla: {lda['best_k']}; c_v={lda['coherence_cv']:.4f}.")
    image("topics_by_medium.png", "Conteos por medio; el volumen del archivo condiciona la comparacion.")
    image("topics_heatmap.png", "Proporciones por medio; -1 indica ausencia de vocabulario util.")
    markdown("## 5. Visualizacion interactiva y reproducibilidad\n\n[pyLDAvis](resultados/lda_visualization.html). Los modelos, diccionario, corpus, parametros, versiones y checksum estan en resultados. La visualizacion HTML puede necesitar red para cargar recursos externos de pyLDAvis.")
    code("display(summary['versions'])\ndisplay(json.loads((RESULTS / 'config.json').read_text()))")
    markdown("## 6. Conclusiones\n\n" + narrative + "\n\n" + event_findings + "\n\n" + format_finding + "\n\nEstos resultados describen vocabulario y agenda de proxies slug. No permiten afirmar que un medio sea mas sesgado, ni comparar framing del mismo acontecimiento sin verificar los pares y recuperar el texto publicado. La siguiente etapa es validar manualmente los casos y extraer una submuestra balanceada.")
    notebook.cells = cells
    nbformat.validate(notebook)
    nbformat.write(notebook, output.parent / "entrega2.ipynb")
    figures = Path(os.path.relpath(output / "figuras", report_dir)).as_posix()
    def figure(filename, caption):
        return "\\begin{figure*}[t]\n\\centering\n\\includegraphics[width=.90\\textwidth]{" + figures + "/" + filename + "}\n\\caption{" + latex_escape(caption) + "}\n\\end{figure*}\n"
    coverage_rows = "\n".join(latex_escape(record.source_id) + " & " + str(record.articulos) + " & " + str(record.con_titulo) + r" \\" for record in coverage.itertuples())
    topic_rows = "\n".join(f"\\item Topico {topic}: " + latex_escape(", ".join(group["palabra"].head(8))) for topic, group in topics.groupby("topic"))
    event_rows = "\n".join(
        latex_escape(medium) + " & " + " & ".join(f"{rates.loc[medium, topic]:.2f}" for topic in ["Reforma tributaria", "Conflicto armado", "Corrupcion"]) + r" \\"
        for medium in rates.index
    )
    latex = r"""\documentclass[10pt,twocolumn]{article}
\usepackage[spanish,es-nodecimaldot]{babel}
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage{geometry,graphicx,booktabs,hyperref,amsmath}
\geometry{top=1.9cm,bottom=2cm,left=1.7cm,right=1.7cm,columnsep=.65cm}
\title{Entrega 2: textometria y topicos del corpus Duque\\Evidencia exploratoria para analizar framing}
\author{Abel Albuez Sanchez, Kelly Joane Leon Torres,\\Juan Camilo Torres Pena, Jesus David Romero Melo\\\small Maestria en Ingenieria de Sistemas y Computacion, Pontificia Universidad Javeriana}
\date{Octubre de 2026}
\begin{document}
\maketitle
\begin{abstract}
""" + latex_escape(narrative) + r""" Los resultados son descriptivos de titulos reconstruidos, no estimaciones validadas de sesgo editorial.
\end{abstract}
\section{Objetivo y datos}
El framing no se reduce a sentimiento ni frecuencia de palabras: requiere comprobar seleccion, enfasis y contexto del mismo acontecimiento \cite{entman}. La entrega estudia candidatos lexicos y tematicos del corpus Duque. Los ocho medios corresponden al inventario unificado, no al corte Duque: aqui hay seis, con cinco series principales y Cambio como muestra tardia.
\begin{center}\begin{tabular}{lrr}\toprule Medio & Articulos & Titulos\\\midrule
""" + coverage_rows + r"""
\bottomrule\end{tabular}\end{center}
El corte de fechas es inclusivo 2018-08-07 a 2022-08-06; no filtra automaticamente precision o asignacion presidencial. No hay cuerpos HTML.
\section{Metodo textometrico}
Se normalizan minusculas y tildes en texto y stopwords; se conservan nombres geograficos y actores. Los n-gramas son secuencias despues de filtrar stopwords y nunca atraviesan documentos. El ajuste rango-frecuencia es descriptivo \cite{zipf}; correlacion alta no demuestra ley de potencia. TTR depende del tamano; MSTTR agrega ventanas completas de 50 tokens. Yule K se calcula como $10^4\sum_w f_w(f_w-1)/N^2$ \cite{yule}.
""" + latex_escape(f"Vocabulario: {lexical['vocabulario']}; tokens: {lexical['tokens']}. Pendiente ajustada: {number(lexical['zipf_slope'])}; R2: {number(lexical['zipf_r2'])}.") + "\n" + figure("zipf_law.png", "Rango-frecuencia: ajuste descriptivo, no certificacion de ley de potencia.") + figure("lexical_diversity.png", "Diversidad lexica por medio con indicadores sensibles a longitud.") + figure("ngrams.png", "N-gramas frecuentes del texto depurado.") + r"""
\section{Temas y acontecimientos candidatos}
Se buscan Reforma tributaria, Conflicto armado y Corrupcion mediante keywords normalizadas y limites de palabra. Se comparan proporciones sobre titulos disponibles, no conteos como sinonimo de sesgo. Las cohortes pueden solaparse. Se revisan ademas ventanas de reforma/paro 2021, bombardeo/menores 2019 y Centros Poblados 2021. Estos casos y sus ejemplos con URLs son candidatos para comprobacion humana; una coincidencia lexica no certifica el mismo hecho. Los contrastes log-odds son descriptivos y no tienen interpretacion causal.
""" + latex_escape(event_findings) + "\n" + r"""
\begin{center}{\scriptsize\begin{tabular}{lrrr}\toprule Medio & Tributaria & Conflicto & Corrupcion\\\midrule
""" + event_rows + r"""
\bottomrule\end{tabular}}\end{center}
Las columnas anteriores son porcentajes de titulos del medio. Cambio (98 titulos) no tiene el mismo alcance temporal; cero coincidencias no prueba ausencia de cobertura.
""" + latex_escape(format_finding) + "\n" + figure("eventos_framing.png", "Cobertura relativa de temas por medio; no es una medicion de framing.") + r"""
\section{Modelado de topicos}
Se usa LDA \cite{lda} sobre todos los documentos con vocabulario util, sin muestreo. Diccionario: frecuencia documental minima 10, maxima .7 y hasta 10000 terminos. Se entrenan K=3,5,7,10,15, diez pasadas, 50 iteraciones internas y semilla 42. Se elige el mayor $c_v$ de entrenamiento, no un optimo universal. La perplejidad aproximada se transforma desde log-bound en base 2. Documentos sin vocabulario reciben -1, no un topico inventado.
""" + latex_escape(f"El mejor K de la grilla fue {lda['best_k']} con c_v={lda['coherence_cv']:.4f}; documentos utiles: {lda['documentos_lda']}.") + "\n\\begin{itemize}\n" + topic_rows + "\n\\end{itemize}\n" + figure("lda_optimization.png", "Comparacion de K: coherencia y perplejidad de entrenamiento.") + figure("topics_by_medium.png", "Topico dominante: conteos condicionados por el volumen del archivo.") + figure("topics_heatmap.png", "Proporciones de topicos por medio.") + "\n\\section{Limitaciones y conclusiones}\n\\begin{itemize}\n" + "\n".join("\\item " + latex_escape(item) for item in LIMITATIONS) + r"""
\end{itemize}
No se concluye que un medio sea mas parcializado. Se obtienen proxies lexicos y de agenda que orientan una submuestra balanceada para recuperar titulares reales/cuerpos, verificar pares de acontecimientos y anotar framing. La comparabilidad requiere controlar fuente de titulo, precision temporal, cobertura de archivo y longitud.
\section{Reproducibilidad}
El notebook integrado carga resultados medidos. Se conservan modelos por K, diccionario, corpus, asignaciones por article\_id, versiones, parametros y SHA-256 del Parquet. No se modifican las bases PostgreSQL ni el export de entrada. Los resultados no sustituyen la validacion manual.
\begin{thebibliography}{9}
\bibitem{entman} R. Entman. Framing: Toward Clarification of a Fractured Paradigm. Journal of Communication, 43(4), 51--58, 1993.
\bibitem{zipf} G. Zipf. Human Behavior and the Principle of Least Effort. Addison-Wesley, 1949.
\bibitem{yule} G. Yule. The Statistical Study of Literary Vocabulary. Cambridge University Press, 1944.
\bibitem{lda} D. Blei, A. Ng y M. Jordan. Latent Dirichlet Allocation. Journal of Machine Learning Research, 3, 993--1022, 2003.
\end{thebibliography}
\end{document}
"""
    (report_dir / "entrega2.tex").write_text(latex, encoding="utf-8")
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    def slide(title, points=(), filename=None):
        page = presentation.slides.add_slide(presentation.slide_layouts[6])
        title_box = page.shapes.add_textbox(Inches(.5), Inches(.3), Inches(12.3), Inches(.8))
        paragraph = title_box.text_frame.paragraphs[0]
        paragraph.text = title
        paragraph.font.name = "Georgia"
        paragraph.font.size = Pt(28)
        paragraph.font.color.rgb = RGBColor(24, 64, 54)
        if filename:
            from PIL import Image
            with Image.open(output / "figuras" / filename) as bitmap:
                width, height = bitmap.size
            scale = min(12.2 / width, 5.1 / height)
            rendered_width, rendered_height = width * scale, height * scale
            page.shapes.add_picture(str(output / "figuras" / filename), Inches((13.333 - rendered_width) / 2), Inches(1.2), width=Inches(rendered_width), height=Inches(rendered_height))
        else:
            body = page.shapes.add_textbox(Inches(.7), Inches(1.4), Inches(11.9), Inches(4.9)).text_frame
            body.word_wrap = True
            for position, point in enumerate(points):
                paragraph = body.paragraphs[0] if position == 0 else body.add_paragraph()
                paragraph.text = point
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(20)
                paragraph.space_after = Pt(16)
        footer = page.shapes.add_textbox(Inches(.5), Inches(6.7), Inches(12.2), Inches(.6)).text_frame.paragraphs[0]
        footer.text = (" | ".join(points) if filename else "Proxies slug; evidencia exploratoria, no prueba de sesgo.")
        footer.font.size = Pt(12)
        return page
    slide("Entrega 2 | Corpus Duque y evidencia exploratoria", [f"{summary['articulos']:,} registros; {summary['medios']} medios presentes", f"{summary['con_titulo']:,} titulos derivados de URL; sin cuerpos HTML", "Textometria, temas/casos candidatos y LDA sobre corpus completo", "Abel Albuez, Kelly Leon, Juan Camilo Torres y Jesus David Romero"])
    slide("Corpus: cinco series principales y Cambio", ["Cambio no constituye una serie completa comparable."], "cobertura_corpus.png")
    slide("Rango-frecuencia: ajuste descriptivo", [f"Pendiente {number(lexical['zipf_slope'])}; R2 {number(lexical['zipf_r2'])}. No demuestra una ley de potencia."], "zipf_law.png")
    slide("Diversidad: controlar longitud y muestra", ["MSTTR agregado usa ventanas completas de 50 tokens."], "lexical_diversity.png")
    slide("N-gramas frecuentes, no significancia estadistica", [f"'{ngram_leader['ngram']}' lidera: {int(ngram_leader['frecuencia']):,}; mezcla de formatos, no sesgo demostrado."], "ngrams.png")
    slide("Temas polemicos: comparar proporciones", [f"La Republica: tributaria {rates.loc['la_republica', 'Reforma tributaria']:.2f}% | RCN: conflicto {rates.loc['noticias_rcn', 'Conflicto armado']:.2f}% | Keywords, no framing validado."], "eventos_framing.png")
    slide("Casos temporales para revision humana", ["Reforma/paro: abril-junio 2021", "Bombardeo/menores: noviembre 2019", "Centros Poblados: agosto-septiembre 2021", "Ejemplos con fechas, medios y URLs disponibles en el notebook.", "Necesitamos validar el mismo hecho y recuperar titulares publicados."])
    slide(f"LDA: mejor K de la grilla = {lda['best_k']}", [f"c_v={lda['coherence_cv']:.4f}; 10 pasadas; semilla 42; evaluacion en entrenamiento."], "lda_optimization.png")
    slide("Topicos por medio: agenda, no sesgo demostrado", ["El topico -1 indica ausencia de vocabulario LDA."], "topics_heatmap.png")
    slide("Conclusiones y siguiente validacion", [f"Vocabulario depurado: {lexical['vocabulario']:,} terminos", f"LDA: {lda['documentos_lda']:,} documentos utiles; K={lda['best_k']}", "No inferir sesgo de frecuencias, sentimiento o topicos aislados.", "Validar fechas, slugs, pares de acontecimientos y balance por medio.", "Extraer una submuestra HTML antes de afirmar diferencias de framing."])
    presentation.save(report_dir / "entrega2.pptx")